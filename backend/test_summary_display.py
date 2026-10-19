import copy
import json
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, os.path.dirname(__file__))

from docx import Document

from pipeline.export_docx import export_docx
from pipeline.export_hwpx import _report_lines
from pipeline.export_markdown import export_markdown
from pipeline.summary_display import get_display_topics

DISPLAY_TOPIC_CASES = json.loads(
    (Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "display-topics.json").read_text(encoding="utf-8")
)


def _result(summary: dict) -> dict:
    return {
        "summary": {"title": "주요 주제 정리", "overview": "요약", "decisions": [], "actions": [], **summary},
        "segments": [],
    }


SECTION_SUMMARY = {
    # Generating topic sections appends their titles to `topics`, so keywords and titles overlap.
    "topics": ["장비 교정", "음향 장비 교정", " 음향 장비 교정 ", "어획량 감소와 수온"],
    "topic_sections": [
        {"topic": "음향 장비 교정", "summary": "교정 본문"},
        {"topic": "어획량 감소와 수온", "summary": "수온 본문"},
        {"topic": "음향 장비 교정 ", "summary": "중복 제목"},
    ],
}


class DisplayTopicsTest(unittest.TestCase):
    def test_shared_display_topic_cases_preserve_source(self) -> None:
        for fixture in DISPLAY_TOPIC_CASES:
            with self.subTest(fixture=fixture["name"]):
                summary = copy.deepcopy(fixture["summary"])
                self.assertEqual(get_display_topics(summary), fixture["expected"])
                self.assertEqual(summary, fixture["summary"])

    def test_prefers_topic_section_titles_and_dedupes(self) -> None:
        self.assertEqual(get_display_topics(SECTION_SUMMARY), ["음향 장비 교정", "어획량 감소와 수온"])

    def test_falls_back_to_summary_topics_without_sections(self) -> None:
        summary = {"topics": ["예산", " 예산", "일정", ""], "topic_sections": []}
        self.assertEqual(get_display_topics(summary), ["예산", "일정"])

    def test_ignores_malformed_values(self) -> None:
        self.assertEqual(get_display_topics(None), [])
        self.assertEqual(get_display_topics({"topics": [None, 3, "예산"], "topic_sections": [None, "x", {"topic": 123}, {"topic": {}}, {"topic": None}]}), ["예산"])


class ExportTopicsTest(unittest.TestCase):
    expected = ["음향 장비 교정", "어획량 감소와 수온"]

    def test_shared_cases_match_in_markdown_full_text_docx_and_hwpx(self) -> None:
        for fixture in DISPLAY_TOPIC_CASES:
            for scope in ("organized", "full"):
                with self.subTest(fixture=fixture["name"], scope=scope), tempfile.TemporaryDirectory() as tmp:
                    result = {**_result(copy.deepcopy(fixture["summary"])), "export_scope": scope}
                    before = copy.deepcopy(result)
                    expected = fixture["expected"]
                    body_texts = [section["summary"] for section in result["summary"]["topic_sections"] if section.get("summary")]
                    for plain in (False, True):
                        path = export_markdown(result, os.path.join(tmp, "out.txt" if plain else "out.md"), plain_text=plain)
                        content = Path(path).read_text(encoding="utf-8")
                        lines = content.split("\n")
                        start = next(index for index, line in enumerate(lines) if line.endswith("주요 주제"))
                        listed = []
                        for line in lines[start + 2:]:
                            if not line.startswith("- "):
                                break
                            listed.append(line[2:])
                        self.assertEqual(listed, expected)
                        for body in body_texts:
                            self.assertIn(body, content)
                    path = export_docx(result, os.path.join(tmp, "out.docx"))
                    document_paragraphs = Document(path).paragraphs
                    paragraphs = [paragraph.text for paragraph in document_paragraphs]
                    start = next(index for index, text in enumerate(paragraphs) if text.endswith("주요 주제"))
                    end = next(index for index, paragraph in enumerate(document_paragraphs)
                               if index > start and paragraph.style.name == "Heading 1")
                    listed = paragraphs[start + 1:end]
                    self.assertEqual(listed, expected)
                    for body in body_texts:
                        self.assertIn(body, "\n".join(paragraphs))
                    lines = _report_lines(result)
                    start, end = lines.index("2. 주요 주제"), lines.index("3. 주제별 내용")
                    self.assertEqual(lines[start + 1:end], [f"- {topic}" for topic in expected] or ["- 없음"])
                    for body in body_texts:
                        self.assertIn(body, "\n".join(lines))
                    self.assertEqual(result, before, "Exports must preserve source topics and every section body")

    def test_markdown_lists_display_topics_once(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = export_markdown(_result(SECTION_SUMMARY), os.path.join(tmp, "out.md"))
            with open(path, encoding="utf-8") as handle:
                content = handle.read()
        block = re.search(r"주요 주제\n\n(.*?)\n\n", content, re.S).group(1)
        self.assertEqual(block.splitlines(), [f"- {topic}" for topic in self.expected])
        self.assertIn("교정 본문", content)

    def test_docx_lists_display_topics_once(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = export_docx(_result(SECTION_SUMMARY), os.path.join(tmp, "out.docx"))
            paragraphs = [paragraph.text for paragraph in Document(path).paragraphs]
        start = next(index for index, text in enumerate(paragraphs) if text.endswith("주요 주제"))
        end = next(index for index, text in enumerate(paragraphs) if index > start and text.endswith("주제별 내용"))
        self.assertEqual(paragraphs[start + 1:end], self.expected)

    def test_hwpx_lists_display_topics_once(self) -> None:
        lines = _report_lines(_result(SECTION_SUMMARY))
        start = lines.index("2. 주요 주제")
        end = lines.index("3. 주제별 내용")
        self.assertEqual(lines[start + 1:end], [f"- {topic}" for topic in self.expected])


if __name__ == "__main__":
    unittest.main()
